from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARCH_PNG = ROOT / "architecture_diagram_publication_aligned.png"
SYNOPSIS_OUT = ROOT / "Agentic_AI_Project_Synopsis_Revised.docx"
EXPLAINER_OUT = ROOT / "Autonomous ML Engineer Crew - Project Explainer Revised.docx"

NAVY = "243F6B"
PALE_BLUE = "EAF1FB"
PALE_GRAY = "F6F7F9"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(80, 80, 80)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_border(cell, color=GRID, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run_font(run, name="Aptos", size=None, bold=None, color=BLACK, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(3)

    for style_name, size, before, after in (
        ("Heading 1", 15, 12, 5),
        ("Heading 2", 12.5, 9, 4),
        ("Heading 3", 11, 7, 3),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.size = Pt(10.5)
        style.font.color.rgb = BLACK
        style.paragraph_format.space_after = Pt(3)


def add_title(doc: Document, title: str, subtitle: str, meta: str | None = None) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(title), "Aptos Display", 23, True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(4)
    set_run_font(p2.add_run(subtitle), "Aptos", 13, False, GRAY, True)
    if meta:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(14)
        set_run_font(p3.add_run(meta), "Aptos", 9.5, False, GRAY)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_paragraph(text, style=f"Heading {level}")


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead:
        set_run_font(p.add_run(bold_lead), bold=True)
    set_run_font(p.add_run(text))


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run_font(p.add_run(item))


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_run_font(p.add_run(item))


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None, font_size: float = 9.3) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)
    for index, text in enumerate(headers):
        cell = header_row.cells[index]
        if widths:
            cell.width = Inches(widths[index])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        set_run_font(paragraph.add_run(text), size=font_size, bold=True, color=RGBColor(255, 255, 255))
    for row_index, row_values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            if widths:
                cell.width = Inches(widths[index])
            set_cell_shading(cell, PALE_BLUE if row_index % 2 else "FFFFFF")
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.03
            set_run_font(paragraph.add_run(str(value)), size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_architecture_figure(doc: Document) -> None:
    if not ARCH_PNG.exists():
        raise FileNotFoundError(f"Architecture image not found: {ARCH_PNG}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(ARCH_PNG), width=Inches(6.78))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(7)
    set_run_font(caption.add_run("Figure 1  Final aligned implementation architecture"), size=9, italic=True, color=GRAY)


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def add_manual_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_together = True
        set_run_font(p.add_run("•  "))
        set_run_font(p.add_run(item))


def build_synopsis() -> None:
    doc = Document()
    configure_document(doc)
    add_title(doc, "Agentic AI Project Synopsis", "Autonomous ML Engineer Crew", "Revised 19 October 2026")

    add_heading(doc, "1 Project Title")
    add_body(doc, "Autonomous ML Engineer Crew")

    add_heading(doc, "2 Team Members")
    add_table(doc, ["Name", "Roll Number"], [
        ["Punya Mahajan", "230611"],
        ["Aniket Mishra", "230648"],
        ["Radhika Singh", "230833"],
    ], [4.6, 2.0], 10)

    add_heading(doc, "3 Problem Definition")
    add_body(doc, "Developing a usable machine-learning prediction service from raw data requires data validation and cleaning, feature preparation, model selection and tuning, held-out evaluation, and deployment. People who have a structured dataset but limited ML expertise or time often cannot complete this workflow. The Autonomous ML Engineer Crew automates these steps through coordinated agents. The user supplies a CSV dataset, names the target column, and describes the prediction goal. The system performs a supported supervised classification or regression workflow. If a trained model does not meet the selected quality threshold, the Evaluation Agent returns structured feedback for another bounded tuning attempt.", "Problem addressed  ")
    add_body(doc, "Intended users include data analysts, small startup teams, and students who have a structured tabular dataset and need a working predictive model without manually building the complete ML pipeline.", "Intended users  ")

    add_heading(doc, "4 Objectives and Expected Outcomes")
    add_body(doc, "The objective and end-user outcome remain unchanged.")
    add_bullets(doc, [
        "Automate data preparation, model selection and tuning, evaluation, and deployment through a coordinated multi-agent pipeline that uses CrewAI and Google ADK.",
        "Implement a self-reflection loop in which the Evaluation Agent applies a task-appropriate metric and justified threshold, then returns actionable feedback to the Model-Selection and Tuning Agent when the model fails.",
        "Deliver a live prediction service with a trained model, a /predict API, a model report card, basic usage and error monitoring, and a reusable pipeline for another compatible dataset.",
    ])
    add_body(doc, "After the user provides a raw CSV dataset, target column, and prediction goal, the system produces a trained model, a live /predict endpoint, a short model report card, usage and error monitoring, and a reusable pipeline. The final workflow result is validated with Pydantic and delivered as JSON plus a human-readable Markdown report card.")

    add_page_break(doc)
    add_heading(doc, "5 Proposed Methodology")
    add_heading(doc, "Agents and Responsibilities", 2)
    add_table(doc, ["Agent", "Framework", "Unchanged responsibility"], [
        ["Data-Prep Agent", "CrewAI", "Validate the CSV and target, profile and clean data, engineer features, and produce a data-quality report."],
        ["Model-Selection and Tuning Agent", "Google ADK", "Train candidate models, run bounded hyperparameter search, log runs, and return model artifacts and metadata."],
        ["Evaluation Agent", "CrewAI", "Run held-out validation, apply task-appropriate metrics, and produce a report card with pass or fail feedback."],
        ["Deployment and Observability Agent", "CrewAI", "Package the accepted model, expose the FastAPI service, and configure runtime monitoring."],
        ["Orchestrator", "CrewAI Flow", "Sequence agents, maintain shared state, validate handoffs, and enforce the bounded evaluation-to-tuning loop."],
    ], [1.75, 1.15, 3.65], 8.8)
    add_body(doc, "CrewAI and Google ADK are the only agentic frameworks. Pydantic, MCP, A2A, MLflow, Langfuse, OpenTelemetry, FastAPI, Docker, and Google Cloud Run support contracts, tools, interoperability, tracing, serving, and deployment; they do not add agents or change agent responsibilities.")

    add_heading(doc, "Framework Boundary and A2A Communication", 2)
    add_body(doc, "CrewAI Flow remains the top-level orchestrator. The existing ADK-based Model-Selection and Tuning Agent is exposed as an A2A service with an Agent Card. The CrewAI Orchestrator calls that service through an A2A client and exchanges typed task messages and artifacts. Each message carries the pipeline run ID, input artifact references, the requested task, the validated result, and any protocol error or timeout. This creates an explicit CrewAI-to-ADK interoperability boundary without adding or renaming agents.")

    add_heading(doc, "Shared State and Pydantic Contracts", 2)
    add_body(doc, "A typed PipelineState is the single source of workflow context. Agents pass artifact references and validated summaries rather than raw dataframe contents. Every handoff uses a Pydantic model, and the orchestrator rejects invalid state transitions before the next agent runs.")
    add_table(doc, ["Contract", "Key content", "Producer and consumer"], [
        ["PipelineState", "run ID, task type, target, status, retry count, artifact references, trace ID", "Orchestrator and all agents"],
        ["DataPrepResult", "dataset hash, schema, quality findings, train and test artifact URIs", "Data-Prep to Model-Selection"],
        ["ModelSelectionResult", "candidate runs, best model URI, parameters, training metrics", "ADK Model-Selection to Evaluation"],
        ["EvaluationResult", "held-out metrics, threshold rationale, pass or fail, retry feedback", "Evaluation to Orchestrator and Model-Selection"],
        ["DeploymentResult", "image digest, service URL, health status, deployment revision", "Deployment to Orchestrator"],
        ["PipelineResult", "final status, model card, endpoint, lineage, trace URL, cost and runtime", "Orchestrator to user"],
    ], [1.35, 3.35, 1.85], 8.6)

    add_page_break(doc)
    add_heading(doc, "MCP Tool Integration", 2)
    add_body(doc, "MCP is mandatory in the implementation. A custom MCP server exposes two tools with JSON schemas and Pydantic-validated returns.")
    add_table(doc, ["MCP tool", "Purpose", "Required evidence"], [
        ["profile_dataset", "Wrap pandas-based schema, missing-value, target, leakage, and dataset-hash checks.", "Successful call, malformed-file failure, validated DatasetProfile response"],
        ["log_and_query_experiment", "Wrap MLflow run creation, metric and parameter logging, artifact registration, and best-run lookup.", "Create and query calls, unavailable-service failure, validated ExperimentRecord response"],
    ], [1.75, 2.85, 1.95], 8.8)

    add_architecture_figure(doc)

    add_heading(doc, "Agentic Workflow and Design Patterns", 2)
    add_numbered(doc, [
        "The Orchestrator validates the request and initializes PipelineState.",
        "The Data-Prep Agent calls the Dataset Profile MCP tool and writes a DataPrepResult.",
        "The Orchestrator sends the prepared-data artifact to the ADK Model-Selection Agent through A2A. The agent trains candidates and calls the MLflow MCP tool.",
        "The Evaluation Agent validates held-out metrics and emits an EvaluationResult.",
        "A passing result routes to deployment. A failing result returns structured feedback for another tuning attempt until the configured retry limit is reached.",
        "The Deployment and Observability Agent deploys the accepted model to Google Cloud Run, runs a health check, and returns a DeploymentResult.",
        "The Orchestrator validates and serializes PipelineResult as JSON and Markdown.",
    ])
    add_body(doc, "The workflow uses sequential orchestration, tool use, conditional routing, an evaluator-optimizer reflection loop, bounded retries, typed state, and explicit terminal success and failure states.")

    add_heading(doc, "Callbacks Monitoring and Error Handling", 2)
    add_body(doc, "The orchestration wrapper defines on-stage-start, on-stage-complete, on-stage-error, retry, and validation-failure callbacks. The ADK agent registers before-agent, after-agent, before-tool, and after-tool callbacks. Every event records the run ID, agent, operation, state transition, duration, result type, retry count, and sanitized error. Callback tests verify normal execution, invalid tool returns, timeouts, and exceptions.")

    add_heading(doc, "Security and Guardrails", 2)
    add_bullets(doc, [
        "Accept only CSV files within configured size and row limits; reject path traversal, missing targets, unsupported task types, and suspicious leakage indicators.",
        "Parse the prediction goal as untrusted input. Detect prompt-injection patterns and abusive or toxic content before it reaches an agent or model.",
        "Validate every A2A message, MCP response, agent handoff, and final output with Pydantic. Invalid data fails closed and creates an audit event.",
        "Keep raw dataset rows inside Python data tools. Send only schema summaries, quality findings, metrics, and artifact references to the LLM.",
        "Use tool allowlists, least-privilege service identities, Secret Manager, log redaction, authenticated endpoints, request limits, and container resource limits.",
    ])

    add_heading(doc, "Observability and Provenance", 2)
    add_body(doc, "Langfuse receives OpenTelemetry traces for the CrewAI Flow, agents, A2A calls, MCP calls, LLM requests, routing decisions, retries, and deployment actions. MLflow remains the system of record for model parameters, metrics, and artifacts. A single run ID links the dataset hash, prepared-data artifact, model run, evaluation, container image digest, deployment revision, and final response. Traces redact raw rows and secrets while retaining timing, token, cost, error, and version metadata.")

    add_heading(doc, "Evaluation Cost and Runtime Plan", 2)
    add_body(doc, "End-to-end evaluation will run on one public regression dataset and one public classification dataset. Repeated runs will cover the success path, one evaluation retry, invalid input, MCP failure, A2A timeout, invalid tool output, exhausted retries, and deployment smoke testing.")
    add_table(doc, ["Area", "Measure", "Evidence"], [
        ["Model quality", "MAE or RMSE for regression; F1 and ROC-AUC for classification; justified pass threshold", "Versioned model report cards and MLflow runs"],
        ["Agent reliability", "task success, schema-valid handoffs, tool failure rate, retry rate, terminal-state correctness", "Automated end-to-end and failure-path tests"],
        ["Runtime", "per-stage duration plus end-to-end p50 and p95 over repeated runs", "Trace export and evaluation table"],
        ["Cost", "LLM input and output tokens, provider price, Cloud Run CPU and memory time, storage and tracing assumptions", "Per-run estimate and monthly scenario"],
        ["Reproducibility", "dataset hash, dependency lock, random seed, model artifact, image digest, deployment revision", "Provenance manifest linked by run ID"],
    ], [1.25, 3.45, 2.0], 8.7)

    add_heading(doc, "Cloud Deployment", 2)
    add_body(doc, "The accepted model is packaged with Docker and served by FastAPI on Google Cloud Run. The service exposes /health and authenticated /predict endpoints. Deployment uses a least-privilege service account and Secret Manager. The final evidence set includes the container image digest, Cloud Run revision, endpoint, health-check result, prediction smoke test, rollback command, and sanitized runtime logs.")

    add_heading(doc, "6 Feasibility and Resources")
    add_body(doc, "The prototype remains limited to small-to-medium structured CSV datasets for supervised classification or regression. Public house-price and customer-churn datasets support repeatable testing, while the demo can accept a compatible user dataset. Candidate models, hyperparameter trials, dataset size, retries, and Cloud Run resources are capped. A standard laptop supports local development, and a modest Cloud Run service supports the final demonstration.")
    add_table(doc, ["Category", "Selected resources"], [
        ["Agentic frameworks", "CrewAI and Google ADK only"],
        ["Contracts and protocols", "Pydantic, MCP, A2A"],
        ["ML and tracking", "pandas, scikit-learn, MLflow"],
        ["Security and tracing", "Callbacks, guardrails, OpenTelemetry, Langfuse"],
        ["Serving and deployment", "FastAPI, Docker, Google Cloud Run, Secret Manager"],
    ], [1.8, 4.75], 9.2)

    add_heading(doc, "7 Team Members and Individual Understanding")
    add_table(doc, ["Member", "Role and responsibility", "Individual understanding"], [
        ["Punya Mahajan 230611", "Model-Selection and Tuning Agent; Orchestrator integration", "Candidate-model selection, bounded tuning, ADK agent and tool logic, CrewAI Flow state, routing, and A2A integration."],
        ["Aniket Mishra 230648", "Data-Prep Agent; Deployment and Observability Agent", "Input validation, data preparation, MCP tool integration, Docker and Cloud Run deployment, API safeguards, and runtime monitoring."],
        ["Radhika Singh 230833", "Evaluation Agent; documentation", "Held-out evaluation, acceptance thresholds, report cards, callback evidence, end-to-end evaluation, and cost and runtime reporting."],
    ], [1.6, 2.1, 2.85], 8.8)
    add_body(doc, "These remain the original ownership areas. All members will test the shared contracts, A2A and MCP integration, guardrails, traces, and final demonstration so that each member can explain and modify the full system.")

    add_heading(doc, "8 Project Timeline")
    add_table(doc, ["Activity", "Week", "Revised deliverable without changing scope"], [
        ["Problem and Design", "September Week 3", "Project charter, fixed agent architecture, typed contracts, A2A and MCP boundaries, threat model"],
        ["Development and Coding", "Late September to October Week 2", "CrewAI Flow, ADK model agent, Pydantic state, two MCP tools, bounded retry"],
        ["Integration and Testing", "October Weeks 3 to 4", "End-to-end workflow, A2A and MCP tests, callbacks, structured logs, failure paths"],
        ["Evaluation and Refinement", "October Week 4 to November Week 1", "Guardrail tests, trace validation, model metrics, latency and cost evaluation"],
        ["Presentation and Interface", "November Weeks 1 to 2", "Minimal upload and query interface, trace walkthrough, demo script"],
        ["Deployment", "November Week 2", "Docker image, Cloud Run service, authenticated /predict, health check, monitoring"],
        ["Final Demo", "Final course week", "Public repository, live pipeline, final report, benchmark evidence, contribution matrix"],
    ], [1.55, 1.5, 3.5], 8.5)

    add_heading(doc, "9 Team Responsibility Matrix")
    add_table(doc, ["Activity", "Punya Mahajan", "Aniket Mishra", "Radhika Singh"], [
        ["Agent Design", "Model selection and orchestration design", "Data preparation and deployment design", "Evaluation and test design"],
        ["Agent and Tool Development", "ADK tools and CrewAI integration", "Data preparation and deployment tools", "Evaluation and report-card tools"],
        ["Integration and Testing", "Workflow and retry integration", "Data pipeline and API integration", "Evaluation tests and sample runs"],
        ["Evaluation and Documentation", "Model and tuning evidence", "Deployment and monitoring evidence", "Metrics, report card, and progress write-up"],
        ["Presentation and Interface", "Architecture and orchestration", "Interface and API demonstration", "Evaluation results and demo data"],
        ["Deployment", "Integration support", "Docker and FastAPI deployment", "Acceptance testing"],
    ], [1.45, 1.7, 1.7, 1.7], 8.2)

    add_page_break(doc)
    add_heading(doc, "10 Minimum Requirements Verification Plan")
    add_body(doc, "The following evidence must exist before the project is presented as compliant. Planning language alone will not count as completion.")
    add_table(doc, ["Requirement", "Implementation commitment", "Completion evidence"], [
        ["1 State sharing", "Typed PipelineState and validated per-agent updates", "State-transition integration test and trace"],
        ["2 MCP tools", "Dataset Profile and MLflow Experiment tools", "Two successful calls plus failure tests"],
        ["3 Pydantic output", "Typed handoffs and PipelineResult JSON and Markdown", "Schema tests and sample serialized output"],
        ["4 Design patterns", "Sequential routing, tools, evaluator-optimizer loop, bounded retry", "Pass, retry, and exhausted-retry traces"],
        ["5 Callbacks", "Stage and ADK agent, model, and tool callbacks", "Correlated lifecycle and error events"],
        ["6 A2A", "CrewAI Orchestrator calls ADK Model-Selection service", "Agent Card, protocol transcript, interoperability test"],
        ["7 Frameworks", "CrewAI and Google ADK only", "Pinned dependencies, imports, and integration run"],
        ["8 Security", "Input, prompt, tool, output, secret, and authorization guardrails", "Adversarial and malformed-response tests"],
        ["9 Observability", "Langfuse and OpenTelemetry plus MLflow lineage", "Trace export and provenance manifest"],
        ["10 Evaluation", "Quality, reliability, runtime, token, and cost evaluation", "Reproducible benchmark and estimate report"],
        ["11 Cloud", "Dockerized FastAPI on Google Cloud Run", "Live endpoint, revision, health check, logs, smoke test"],
    ], [1.25, 3.25, 2.05], 8.1)

    doc.save(SYNOPSIS_OUT)


def build_explainer() -> None:
    doc = Document()
    configure_document(doc)
    add_title(doc, "Autonomous ML Engineer Crew", "Project Ideation and Implementation Plan", "Revised 19 October 2026")

    add_heading(doc, "Overview")
    add_body(doc, "The Autonomous ML Engineer Crew automates the supported machine-learning workflow from a raw CSV dataset to a deployed prediction service. A user provides a dataset, target column, and prediction goal. The same five established agents validate and prepare the data, train candidates, evaluate the selected model, retry tuning when necessary, and deploy an accepted model behind a /predict API.")
    add_body(doc, "The revision does not change the objective, intended users, agent roster, agent responsibilities, supervised classification and regression scope, or evaluator-to-tuner retry loop. It defines the technical contracts and evidence needed to satisfy the course minimum requirements.")

    add_heading(doc, "User Outcome")
    add_bullets(doc, [
        "A trained model that predicts new tabular records.",
        "A live authenticated /predict endpoint and a /health endpoint.",
        "A concise model report card describing the selected model, evaluation metric, threshold, and result.",
        "Usage and error monitoring linked to the execution trace.",
        "A reusable pipeline for another compatible CSV dataset.",
    ])

    add_heading(doc, "Fixed Agent Architecture")
    add_table(doc, ["Agent", "Responsibility", "Implementation framework"], [
        ["Data-Prep Agent", "Clean and validate data, handle missing values, engineer features, flag quality issues", "CrewAI"],
        ["Model-Selection and Tuning Agent", "Train candidate models and run bounded hyperparameter search", "Google ADK"],
        ["Evaluation Agent", "Evaluate held-out data, apply a justified threshold, issue pass or retry feedback", "CrewAI"],
        ["Deployment and Observability Agent", "Package, deploy, and monitor the accepted model API", "CrewAI"],
        ["Orchestrator", "Manage order, shared state, validation, routing, and retry termination", "CrewAI Flow"],
    ], [1.65, 3.55, 1.35], 8.8)
    add_body(doc, "CrewAI and Google ADK are the only agentic frameworks. The mapping makes the existing ADK model-selection core the remote specialist while CrewAI continues to coordinate the surrounding workflow.")

    add_page_break(doc)
    add_architecture_figure(doc)

    add_heading(doc, "Implementation Contracts")
    add_heading(doc, "Shared State", 2)
    add_body(doc, "CrewAI Flow owns a Pydantic PipelineState containing the run ID, dataset and artifact references, task type, target column, current stage, retry count, model and evaluation metadata, deployment result, trace ID, and sanitized errors. Each agent reads the state it needs and writes a validated result for the next stage.")

    add_heading(doc, "Structured Outputs", 2)
    add_body(doc, "Pydantic models validate DataPrepResult, ModelSelectionResult, EvaluationResult, DeploymentResult, and the terminal PipelineResult. The final result is serialized to JSON for machines and rendered as a Markdown model report card for people. Invalid fields stop the workflow and create an audit event.")

    add_heading(doc, "MCP Tools", 2)
    add_body(doc, "MCP is a required integration rather than an optional extension. The Dataset Profile tool performs schema, missingness, target, leakage, and dataset-hash checks. The MLflow Experiment tool creates and queries runs, metrics, parameters, and artifacts. Both tools publish schemas, return Pydantic-validated data, and have success and failure tests.")

    add_heading(doc, "A2A Interoperability", 2)
    add_body(doc, "The ADK Model-Selection and Tuning Agent runs as an A2A service with an Agent Card. The CrewAI Orchestrator invokes it through an A2A client. Requests and results carry the run ID and artifact references so the same trace and provenance chain survives the framework boundary.")

    add_heading(doc, "Agentic Patterns", 2)
    add_body(doc, "The design uses a sequential workflow for the normal path, tool use for dataset and experiment operations, conditional routing at the evaluation gate, and an evaluator-optimizer loop for retuning. A configured retry budget prevents an endless loop and routes exhausted attempts to a clear terminal failure.")

    add_heading(doc, "Monitoring Security and Traceability")
    add_heading(doc, "Callbacks and Logging", 2)
    add_body(doc, "The orchestrator wraps every stage with start, complete, validation-failure, retry, and error callbacks. The ADK agent registers before-agent, after-agent, before-tool, and after-tool callbacks. Each event includes the run ID, stage, agent, state transition, duration, result type, retry count, and sanitized error. These events support debugging without logging raw dataset rows.")

    add_heading(doc, "Guardrails", 2)
    add_manual_bullets(doc, [
        "Validate file type, size, target column, task type, path, schema, and leakage indicators before agent execution.",
        "Treat the prediction goal and dataset text as untrusted input; detect prompt injection and abusive or toxic language.",
        "Validate A2A messages, MCP responses, handoffs, and final output with Pydantic and fail closed on invalid data.",
        "Restrict tools and service accounts to the permissions they need; keep secrets in Secret Manager and redact them from traces.",
        "Keep raw rows in deterministic Python tools and send only summaries, metrics, and artifact references to the LLM.",
    ])

    add_heading(doc, "Observability and Provenance", 2)
    add_body(doc, "Langfuse receives OpenTelemetry spans from CrewAI, ADK, A2A, MCP, LLM, evaluation, and deployment operations. MLflow records model parameters, metrics, and artifacts. One run ID links the dataset hash, preprocessing artifact, model run, evaluation decision, image digest, Cloud Run revision, and final response. This supplies an auditable path from input to deployment.")

    add_heading(doc, "Evaluation and Cost Evidence", 2)
    add_body(doc, "The final evaluation uses a public regression dataset and a public classification dataset. Repeated end-to-end runs measure model quality, pipeline success, schema-valid handoffs, retries, failures, per-stage duration, p50 and p95 runtime, LLM token use, and cloud compute time. The cost report multiplies measured usage by documented provider and Cloud Run unit prices and includes a per-run estimate plus a small monthly usage scenario.")

    add_heading(doc, "Cloud Deployment", 2)
    add_body(doc, "Docker packages the accepted model with a FastAPI service. Google Cloud Run hosts /health and authenticated /predict endpoints using a least-privilege service account and Secret Manager. Completion evidence includes the image digest, deployed revision, endpoint, health check, prediction smoke test, rollback command, sanitized logs, and a linked Langfuse trace.")

    add_page_break(doc)
    add_heading(doc, "End to End Workflow")
    add_numbered(doc, [
        "Validate the uploaded CSV, target column, and prediction goal, then initialize typed PipelineState.",
        "Run the Data-Prep Agent and the Dataset Profile MCP tool. Store the cleaned-data artifact and DataPrepResult.",
        "Call the ADK Model-Selection Agent over A2A. Train candidates and record runs through the MLflow MCP tool.",
        "Run the Evaluation Agent. If the model fails, return typed feedback for another bounded tuning attempt.",
        "When the model passes, run the Deployment and Observability Agent to build and deploy the service on Cloud Run.",
        "Run health and prediction smoke tests, finalize the trace and provenance manifest, and return PipelineResult as JSON and Markdown.",
    ])

    add_heading(doc, "Definition of Done")
    add_table(doc, ["Requirement", "The project is complete when"], [
        ["State", "A trace and integration test show each agent consuming the previous agent's typed output."],
        ["MCP", "Two MCP tools run successfully in the workflow and their failure paths are tested."],
        ["Structured output", "Pydantic validates every handoff and the final JSON and Markdown result."],
        ["Patterns", "Pass, retry, and exhausted-retry executions follow the designed routing."],
        ["Callbacks", "Intermediate output, lifecycle, validation, retry, and error events are recorded."],
        ["A2A", "A CrewAI participant and the ADK model-selection service exchange real A2A messages."],
        ["Frameworks", "Pinned dependencies and runtime evidence show CrewAI and ADK, with no third agentic framework."],
        ["Security", "Injection, abusive input, malformed tool data, unauthorized action, and secret-leak tests pass."],
        ["Observability", "A trace connects agents, A2A, MCP, model lineage, and the deployed revision."],
        ["Evaluation", "The benchmark report includes model quality, reliability, runtime, tokens, and cost."],
        ["Cloud", "Cloud Run serves the model and the health and prediction smoke tests pass."],
    ], [1.55, 5.05], 8.7)

    add_heading(doc, "Scope Boundaries")
    add_table(doc, ["Boundary"], [
        ["Supported tasks: supervised classification and regression on structured CSV datasets only."],
        ["Resource limits: cap dataset size, candidate models, hyperparameter trials, retries, and cloud resources."],
        ["Fixed architecture: keep the five established agents and their original responsibilities."],
        ["Framework limit: use no additional agentic framework; supporting libraries and protocols remain infrastructure."],
        ["Evidence rule: do not claim completion until the required tests, traces, reports, and deployment evidence exist."],
    ], [6.55], 8.3)

    doc.save(EXPLAINER_OUT)


if __name__ == "__main__":
    build_synopsis()
    build_explainer()
    print(SYNOPSIS_OUT)
    print(EXPLAINER_OUT)
